from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT = Path(r"D:\ЯндексДиск\Yandex.Disk\ПРОЕКТЫ\КОРА_35_ЮБИЛЕЙ")
REPORT_DIR = PROJECT / "06_НАПРАВЛЕНИЯ" / "Дети" / "экскурсия" / "Бухгалтерия детс экскурсии" / "ОТЧЕТНЫЙ_ПАКЕТ_2026-07-20"
REPORT_OUT = REPORT_DIR / "ОТЧЕТ_ДЕТСКАЯ_ЭКСКУРСИЯ_ЗАТРАТЫ_И_ПРЕМИИ_2026-07-20.docx"
RECEIPT_DIR = PROJECT / "02_ОПЕРАЦИОНКА" / "ДОКУМЕНТООБОРОТ"
RECEIPT_OUT = RECEIPT_DIR / "РАСПИСКА_ПЕРЕДАЧА_ЗАДАТКА_ЛИТЕЙЩИК_5000_2026-07-20.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 46, 38)
MUTED = RGBColor(93, 105, 96)
LIGHT = "F2F4F7"
CAUTION = "FFF4C2"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def set_font(run, size=11, bold=None, color=INK, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_document(doc, running_label):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_font(header.add_run(running_label), size=9, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    set_font(footer.add_run("КОРА 35 · 20.07.2026 · "), size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)


def add_title(doc, title, subtitle, metadata, title_size=23, title_before=16, subtitle_after=16, metadata_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(title_before)
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(title), size=title_size, bold=True, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(subtitle_after)
    set_font(p.add_run(subtitle), size=13, color=MUTED)
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(metadata_after)
        set_font(p.add_run(f"{label}: "), bold=True)
        set_font(p.add_run(value))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=60, start=120, bottom=60, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_table(doc, headers, rows, widths, numeric_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if index in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(text), size=9.5, bold=True, color=INK)
    for row_data in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row_data):
            p = cells[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if index in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(text)), size=9.5, bold=str(text).startswith("Итого") or str(text).startswith("42 700"))
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, text, fill=CAUTION):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_cell_shading(table.cell(0, 0), fill)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(text), bold=True)
    set_table_geometry(table, [CONTENT_DXA])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_report():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc, "Управленческий отчёт · детская экскурсия")
    doc.styles["List Number"].paragraph_format.space_after = Pt(4)
    add_title(
        doc,
        "ОТЧЁТ ПО ЗАТРАТАМ",
        "Детская экскурсия КОРА 35 · прямые расходы и расчётный вклад организаторов",
        [("Дата мероприятия", "24 июня 2026 г."), ("Дата отчёта", "20 июля 2026 г."), ("Статус", "для сверки и бухгалтерского закрытия")],
    )
    add_callout(doc, "Прямой факт: 26 950,59 ₽. С расчётной оценкой труда организаторов: 42 700,59 ₽. Остаток до потолка детского блока 48 000 ₽ — 5 299,41 ₽.")

    doc.add_heading("1. Итоговая рамка", level=1)
    add_table(doc, ["Показатель", "Сумма", "Статус"], [
        ["Подтверждённые прямые расходы", "26 950,59 ₽", "рабочий подтверждённый факт"],
        ["Расчётный эквивалент труда", "15 750,00 ₽", "не утверждено к выплате"],
        ["Полная управленческая стоимость", "42 700,59 ₽", "прямой факт + оценка труда"],
        ["Потолок детского блока", "48 000,00 ₽", "утверждённая рамка"],
        ["Остаток до потолка", "5 299,41 ₽", "при отсутствии новых расходов"],
    ], [3400, 1800, 4160], numeric_cols=(1,))
    p = doc.add_paragraph()
    set_font(p.add_run("Доплата 1 054 ₽: "), bold=True)
    set_font(p.add_run("передана повару наличными 20.07.2026, закрыла ранее учтённый остаток по столовой и не прибавляется к итогу повторно."))

    doc.add_heading("2. Прямые расходы", level=1)
    add_table(doc, ["Статья", "Сумма", "Подтверждение / статус"], [
        ["Столовая: 11 взрослых и 6 детей", "2 550,00 ₽", "1 496 ₽ зачтено; 1 054 ₽ доплачено 20.07"],
        ["Копировальные услуги", "366,00 ₽", "товарный чек"],
        ["Награды за конкурсы", "125,00 ₽", "чек"],
        ["Бананы", "749,00 ₽", "чек"],
        ["Сувенирные пакетики", "600,00 ₽", "чек"],
        ["Пластилин", "1 190,00 ₽", "чек"],
        ["Автобус НефАЗ 5299", "10 000,00 ₽", "наличными по акту"],
        ["Такси Яндекс Go", "112,00 ₽", "PDF-отчёт"],
        ["Реквизит Д. Кузьмичева", "2 754,55 ₽", "фото и PDF-чеки"],
        ["Закупки Прусакова В.В.", "8 504,04 ₽", "фото чеков"],
        ["Итого прямых расходов", "26 950,59 ₽", ""],
    ], [3500, 1700, 4160], numeric_cols=(1,))
    p = doc.add_paragraph()
    set_font(p.add_run("Техническое расхождение: "), bold=True)
    set_font(p.add_run("чеки Прусакова дают 8 504,04 ₽; расчёт аванса 10 000 − 1 496 даёт 8 504 ₽. Разница 0,04 ₽."))

    doc.add_heading("3. Расчётный вклад организаторов", level=1)
    p = doc.add_paragraph("Рабочая ставка — 450 ₽/час. Оценка нужна для управленческого решения и не является автоматическим начислением.")
    add_table(doc, ["Участник", "Часы", "Эквивалент"], [
        ["Дмитрий Коган", "8", "3 600 ₽"],
        ["Дмитрий Кузьмич", "7", "3 150 ₽"],
        ["Нияз Кашапов", "5", "2 250 ₽"],
        ["Евгения Сенина", "6", "2 700 ₽"],
        ["Настя Бурд", "4", "1 800 ₽"],
        ["С. Б. Касьянов", "3", "1 350 ₽"],
        ["Никита (брелоки)", "2", "900 ₽"],
        ["Итого", "35", "15 750 ₽"],
    ], [5000, 1360, 3000], numeric_cols=(1, 2))

    doc.add_heading("4. Артефакты и закрытие", level=1)
    for text in [
        "В исходной бухгалтерской папке проинвентаризировано 20 файлов: смета-факт DOCX, PDF Яндекс Go, фотографии чеков и PDF-чеки участников.",
        "Полный перечень сохранён рядом с отчётом в реестре артефактов от 20.07.2026.",
        "Служебный файл _tmp_partner_card_page1.png не считается финансовым подтверждением без отдельной идентификации.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("5. Что требуется до бухгалтерского закрытия", level=1)
    for text in [
        "Приложить подтверждение получения поваром 1 054 ₽, если оно оформлено.",
        "Проверить наличие подписанного акта по автобусу в финальном пакете.",
        "Подтвердить отсутствие дополнительных чеков.",
        "Утвердить или скорректировать часы и премии организаторам.",
        "Проверить в бухгалтерии форму документов и основание для выплат.",
    ]:
        doc.add_paragraph(text, style="List Number")
    doc.save(REPORT_OUT)


def add_form_line(doc, label, line="____________________________________________________________"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(f"{label}: "), bold=True)
    set_font(p.add_run(line))


def build_receipt():
    RECEIPT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc, "Документооборот · задаток базе «Литейщик»")
    section = doc.sections[0]
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.58)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.0
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(8)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(4)
    doc.styles["List Bullet"].font.size = Pt(10.5)
    doc.styles["List Bullet"].paragraph_format.space_after = Pt(3)
    doc.styles["List Bullet"].paragraph_format.line_spacing = 1.0
    add_title(
        doc,
        "РАСПИСКА О ПОЛУЧЕНИИ ЗАДАТКА",
        "Бронирование базы «Литейщик» для корпоративного выезда КОРА 35",
        [("Дата события", "15 августа 2026 г.; сотрудники 10:00–15:00, организаторы с 08:30"), ("Сумма", "5 000 (Пять тысяч) рублей 00 копеек")],
        title_size=20,
        title_before=6,
        subtitle_after=8,
        metadata_after=0,
    )
    add_form_line(doc, "Город")
    add_form_line(doc, "Дата", "«___» __________ 2026 г.")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    set_font(p.add_run("Я, "), bold=True)
    set_font(p.add_run("________________________________________________________________________,"))
    add_form_line(doc, "Паспорт / документ", "серия ______ № __________, выдан ________________________________")
    add_form_line(doc, "Действую от имени")
    add_form_line(doc, "Основание полномочий")
    p = doc.add_paragraph()
    set_font(p.add_run("получил(а) от "), bold=True)
    set_font(p.add_run("__________________________________________________________________"))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run("денежные средства в размере 5 000 (Пять тысяч) рублей 00 копеек "), bold=True)
    set_font(p.add_run("в качестве задатка / обеспечительного платежа за бронирование базы отдыха «Литейщик» для проведения корпоративного выезда КОРА 35 15 августа 2026 года."))
    add_form_line(doc, "Полная стоимость аренды / услуги")
    add_form_line(doc, "Остаток к оплате и срок")
    doc.add_heading("Условия задатка", level=1)
    for text in [
        "Полученная сумма засчитывается в окончательную стоимость: да / нет.",
        "Возврат при отмене со стороны площадки: ______________________________________________.",
        "Возврат при отмене со стороны заказчика: _____________________________________________.",
        "Иные условия: _____________________________________________________________________.",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    add_form_line(doc, "Договор / приложение / переписка")
    add_form_line(doc, "Реквизиты подтверждающего документа")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    add_form_line(doc, "Деньги получил(а)", "__________________ / __________________________________________")
    add_form_line(doc, "Деньги передал(а)", "__________________ / __________________________________________")
    add_form_line(doc, "Контакт получателя")
    add_callout(doc, "Перед передачей денег бухгалтерии следует проверить статус получателя и название платежа. Если деньги принимает организация или ИП, предпочтительны договор, счёт и кассовый или банковский документ; одной расписки может быть недостаточно для бухгалтерского закрытия.")
    doc.save(RECEIPT_OUT)


if __name__ == "__main__":
    build_report()
    build_receipt()
    print(REPORT_OUT)
    print(RECEIPT_OUT)
