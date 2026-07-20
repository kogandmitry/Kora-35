from pathlib import Path
from shutil import copy2

from docx import Document


PROJECT = Path(r"D:\ЯндексДиск\Yandex.Disk\ПРОЕКТЫ\КОРА_35_ЮБИЛЕЙ")
ARCHIVE = PROJECT / "_АРХИВ_ВЕРСИЙ" / "date-time-budget-sync-2026-07-20"
RECEIPT = PROJECT / "02_ОПЕРАЦИОНКА" / "ДОКУМЕНТООБОРОТ" / "РАСПИСКА_ПЕРЕДАЧА_ЗАДАТКА_ЛИТЕЙЩИК_5000_2026-07-20.docx"
CONCEPT = PROJECT / "04_ЗАКАЗЧИКАМ" / "АКТУАЛЬНЫЙ_ПАКЕТ_ДОКУМЕНТОВ" / "КОРА35_ЛИТЕЙЩИК_КОНЦЕПЦИЯ_ПРОГРАММА_СМЕТА_2026-07-14.docx"


def all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def replace_in_runs(document, replacements):
    for paragraph in all_paragraphs(document):
        for run in paragraph.runs:
            for old, new in replacements:
                if old in run.text:
                    run.text = run.text.replace(old, new)


def set_row(table, label, values):
    for row in table.rows:
        if row.cells and row.cells[0].text.strip() == label:
            for index, value in enumerate(values):
                row.cells[index].text = value
            return True
    return False


def backup(path):
    ARCHIVE.mkdir(parents=True, exist_ok=True)
    target = ARCHIVE / path.name
    if not target.exists():
        copy2(path, target)


def update_receipt():
    backup(RECEIPT)
    document = Document(RECEIPT)
    replace_in_runs(
        document,
        [
            ("16 августа 2026 г.", "15 августа 2026 г.; сотрудники 10:00–15:00, организаторы с 08:30"),
            ("16 августа 2026 года", "15 августа 2026 года"),
        ],
    )
    document.save(RECEIPT)


def update_concept():
    backup(CONCEPT)
    document = Document(CONCEPT)
    replace_in_runs(
        document,
        [
            ("16 августа", "15 августа"),
            ("воскресенье, 15 августа", "суббота, 15 августа"),
            ("15 августа 2026 года, воскресенье", "15 августа 2026 года, суббота"),
            ("воскресенье 15 августа", "субботу 15 августа"),
            ("время доступа команды до 10:00", "время доступа организаторов с 08:30"),
            ("08:45–10:00", "08:30–10:00"),
            ("выезд в воскресенье", "выезд в субботу"),
        ],
    )

    date_paragraph = document.paragraphs[3]
    date_paragraph.text = "Дата события: суббота, 15 августа 2026 года · сотрудники 10:00–15:00 · организаторы с 08:30"

    prep_table = document.tables[0]
    prep_table.rows[1].cells[0].text = "08:30"
    prep_table.rows[1].cells[1].text = "прибытие организаторов, руководителей зон, технической команды и клининга"

    program_table = document.tables[1]
    for row in program_table.rows:
        if row.cells[0].text.strip() == "14:00–14:30":
            row.cells[2].text = "DJ/танцы, викторина, финальные мастер-классы, тихие беседки, фото и подача торта около 14:20"

    summary_table = document.tables[2]
    set_row(summary_table, "Площадка, клининг, безопасность и операционные мелочи", ["Площадка и операционка", "119 900 ₽"])
    set_row(summary_table, "Питание", ["Питание", "191 500 ₽"])
    set_row(summary_table, "Программа и взрослые активности", ["Программа и активности Нияза", "76 000 ₽"])
    set_row(summary_table, "Детский контур", ["Детский контур", "41 000 ₽"])
    set_row(summary_table, "Признание сотрудников", ["Подарки", "105 000 ₽"])
    set_row(summary_table, "Медиа и цифровой контур", ["Медиа и цифровой контур", "270 000 ₽"])
    set_row(summary_table, "Управление, помощники, подготовка и постсобытийная работа", ["Управление и после", "117 000 ₽"])
    set_row(summary_table, "ИТОГО", ["ИТОГО полного локального контура", "1 050 400 ₽"])
    set_row(summary_table, "Свободно до потолка", ["Превышение потолка полного контура", "50 400 ₽"])

    detail_table = document.tables[3]
    set_row(
        detail_table,
        "DJ, звук, микрофоны и показ видео",
        ["LED-экран и музыкальное оборудование", "пакет", "85 000", "сумма задана пользователем", "единый пакет: экран, музыка, микрофоны, подключение и резервный носитель"],
    )
    set_row(
        detail_table,
        "Фотограф 15 августа",
        ["Фото- и видеофиксация 15 августа", "пакет", "20 000", "20 000–45 000", "единая строка: ключевые блоки, репортаж и обработка"],
    )
    set_row(
        detail_table,
        "Видеофиксация 15 августа",
        ["Инфраструктура управления: Claude и ChatGPT", "май–август", "12 000", "12 000", "закрытый управленческий контур; сохранить подтверждения оплат"],
    )
    set_row(detail_table, "ИТОГО", ["ИТОГО полного локального контура", "", "1 050 400", "", ""])

    options_table = document.tables[4]
    set_row(
        options_table,
        "Снеки на выезд",
        ["Торт в конце мероприятия", "15 000", "заменяет снеки без увеличения рабочего сценария; подача около 14:20"],
    )
    document.save(CONCEPT)


if __name__ == "__main__":
    update_receipt()
    update_concept()
    print(RECEIPT)
    print(CONCEPT)
